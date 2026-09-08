from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "reports"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "NBA_MVP_CRISP_DM_Final_Report.docx"

NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
ORANGE = "D97706"
GREEN = "2F855A"
GRAY = "5F6874"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "E8EEF5"
WHITE = "FFFFFF"
BLACK = "202124"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=None, color=BLACK, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9, color=GRAY)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        "Title": (30, NAVY, 0, 8),
        "Subtitle": (15, DARK_BLUE, 0, 8),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in specs.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)


def add_running_furniture(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("NBA MVP VOTE-SHARE MODEL  |  CRISP-DM FINAL REPORT")
    set_run_font(r, size=8.5, color=GRAY, bold=True)
    p.paragraph_format.space_after = Pt(0)
    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(item)
        set_run_font(r)


def add_callout(doc, label, text, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    set_repeat_table_header(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    lr = p.add_run(f"{label}: ")
    set_run_font(lr, bold=True, color=accent)
    tr = p.add_run(text)
    set_run_font(tr, color=BLACK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths_dxa, numeric_columns=None):
    numeric_columns = set(numeric_columns or [])
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, size=9, color=WHITE, bold=True)
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            if row_number % 2:
                set_cell_shading(cells[index], LIGHT_GRAY)
            p = cells[index].paragraphs[0]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if index in numeric_columns
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=9)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc, path, caption, width=6.3):
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("title", caption)
    shape._inline.docPr.set("descr", caption)
    cp = doc.add_paragraph(caption, style="Caption")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER


doc = Document()
configure_styles(doc)
for section in doc.sections:
    configure_page(section)
    add_running_furniture(section)

# Cover: editorial-cover pattern with a restrained analytical palette.
for _ in range(4):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("CRISP-DM DATA SCIENCE REPORT")
set_run_font(r, size=10, color=ORANGE, bold=True)
p.paragraph_format.space_after = Pt(18)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Predicting NBA MVP Vote Share")
set_run_font(r, size=30, color=NAVY, bold=True)

p = doc.add_paragraph(style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Regression, season-level ranking, and era-aware evaluation using player-season data from 1982–2022")
set_run_font(r, size=14, color=DARK_BLUE)

doc.add_paragraph()
add_callout(
    doc,
    "Central finding",
    "Regular-season player and team statistics provide useful predictive signal. The locked Histogram Gradient Boosting model achieved 0.0263 all-player RMSE, 0.8934 NDCG@5, and ranked the actual MVP first in three of four final-test seasons.",
    fill="F7F2E8",
    accent=ORANGE,
)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(34)
r = p.add_run("Final analytical report  |  September 2026")
set_run_font(r, size=10.5, color=GRAY, italic=True)

doc.add_page_break()

doc.add_heading("Executive summary", level=1)
add_body(doc, "This project evaluated whether regular-season performance can predict the numerical MVP voting target, award_share, and rank leading candidates within each season. The analysis treated MVP prediction as both a regression problem and a season-level information-retrieval problem because most players receive no votes and the practical question concerns ordering the small group of serious candidates.")
add_body(doc, "The final design used chronological development through 2018 and an untouched 2019–2022 test. Histogram Gradient Boosting was locked as the official model before the final test, with an Extra Trees hurdle model retained as the principal structural challenger. The champion passed all six prespecified success criteria.")
add_table(
    doc,
    ["Final-test measure", "Histogram GB", "Interpretation"],
    [
        ["All-player RMSE", "0.0263", "Best numerical error"],
        ["Vote-recipient RMSE", "0.1553", "Best recipient error"],
        ["R²", "0.8033", "Strong explained variation"],
        ["NDCG@5", "0.8934", "Strong top-candidate ordering"],
        ["Actual MVP mean rank", "1.25", "Winner never below second"],
        ["Top-1 accuracy", "75%", "Three of four winners"],
    ],
    [3000, 1800, 4560],
    numeric_columns={1},
)
add_callout(doc, "Decision", "Use Histogram GB for numerical vote-share forecasts and require the Extra Trees hurdle ranking as a calibration and disagreement check.", fill=PALE_BLUE)

doc.add_page_break()
doc.add_heading("1. Business understanding", level=1)
doc.add_heading("1.1 Research question", level=2)
add_body(doc, "Can regular-season player statistics and team performance predict MVP vote share and correctly rank the leading MVP candidates within each season?")
doc.add_heading("1.2 Analytical framing", level=2)
add_bullets(doc, [
    "Regression: predict the continuous award_share target.",
    "Ranking: order candidates separately within each season.",
    "Chronological generalization: train on the past and evaluate on later seasons.",
    "Decision support: produce forecasts and candidate rankings, not causal judgments about who deserves the award.",
])
doc.add_heading("1.3 Success criteria", level=2)
add_body(doc, "The champion was required to outperform the Win Shares baseline on all-player RMSE, recipient RMSE, NDCG@5, and season-total calibration; achieve a mean winner rank no worse than third; and obtain at least 50% top-1 accuracy. All six criteria passed on the locked test.")

doc.add_heading("2. Data understanding", level=1)
add_body(doc, "The uploaded dataset contains 17,697 player-team-season rows, 55 columns, and 41 seasons from 1982 through 2022. The target is extremely zero-heavy: meaningful MVP support is concentrated among a small number of candidates each season. Repeated names are expected because players appear across seasons and traded players may have team-specific and combined TOT records.")
add_table(
    doc,
    ["Audit item", "Observed result", "Treatment"],
    [
        ["Temporal coverage", "1982–2022", "Preserve chronological order"],
        ["Rows / columns", "17,697 / 55", "Retain raw source identity"],
        ["TOT rows", "1,954", "Retain; mark is_tot"],
        ["Repeated season-player rows", "44", "Investigate, do not auto-delete"],
        ["One repeated season-player-team key", "1989 Charles Jones, WSB", "Retain as audited exception"],
        ["Target distribution", "Large zero mass", "Use weighted regression and ranking metrics"],
    ],
    [2550, 2400, 4410],
)
add_figure(doc, ROOT / "reports" / "figures" / "target_distribution.png", "Figure 1. The MVP target is dominated by zero-share player-seasons, motivating candidate-sensitive metrics.")

doc.add_heading("3. Data preparation", level=1)
add_body(doc, "Preparation decisions were investigated before modification. The analysis retained repeated players, TOT records, and statistical outliers because these can represent legitimate player-season structure or historically exceptional performance.")
add_table(
    doc,
    ["Issue", "Final decision", "Rationale"],
    [
        ["Minimum minutes", "mp ≥ 100", "Broad coverage; all recipients retained"],
        ["Sensitivity threshold", "mp ≥ 500", "Tests robustness to fringe players"],
        ["TOT team context", "Team variables set missing", "Combined row has no single team context"],
        ["Missing shooting percentages", "Zero plus indicator", "Represents no attempts without hiding structure"],
        ["Remaining numerical missingness", "Training median", "Prevents future-information leakage"],
        ["Outliers", "Retain", "Extreme excellence is central to MVP prediction"],
        ["Position", "Primary position; one-hot", "Stable categorical representation"],
    ],
    [2350, 2450, 4560],
)

doc.add_heading("4. Era-aware feature engineering", level=1)
add_body(doc, "Raw NBA statistics are not directly comparable across eras because pace, role specialization, shooting efficiency, schedule structure, and league strategy change. The final pipeline therefore combines raw measurements with within-season percentile ranks.")
add_body(doc, "Season percentiles were generated for points, minutes, true shooting, PER, Win Shares, WS/48, BPM, VORP, and team winning percentage. The target was removed before these features were computed, preventing target leakage.")
add_figure(doc, ROOT / "reports" / "figures" / "statistical_environment_by_season.png", "Figure 2. Selected league statistics change materially over time, supporting era-aware feature engineering.")

doc.add_heading("5. Validation design and leakage control", level=1)
add_body(doc, "All modeling decisions were made chronologically. Development used seasons through 2014, validation used 2015–2018, and final testing used 2019–2022. The locked final model was not altered after test outcomes were viewed.")
add_bullets(doc, [
    "Player identity was excluded to prevent reputation memorization.",
    "Calendar season was excluded; era information entered through within-season comparisons.",
    "Imputation and categorical encoding were fitted on training data only.",
    "award_share was excluded from every predictor transformation.",
    "No model or hyperparameter was changed after the final test was opened.",
])
add_figure(doc, ROOT / "reports" / "figures" / "chronological_evaluation_design.png", "Figure 3. Chronological separation used for development, validation, and final testing.")

doc.add_heading("6. Modeling", level=1)
add_body(doc, "The candidate set progressed from trivial and domain baselines to regularized linear models, nonlinear ensembles, and explicit hurdle structures. Cluster assignments were evaluated but excluded from the final predictor because they did not provide stable incremental value.")
add_table(
    doc,
    ["Model", "Role", "Benefit", "Main risk"],
    [
        ["Zero / historical mean", "Trivial baselines", "Sanity checks", "No useful ranking"],
        ["WS OLS", "Domain baseline", "Transparent", "Weak numerical calibration"],
        ["Weighted Ridge", "Ranking comparator", "Interpretable", "Overallocates vote share"],
        ["Histogram GB", "Official model", "Best RMSE and R²", "High zero-boundary clipping"],
        ["Extra Trees hurdle", "Required challenger", "Best NDCG and totals", "More complex; lower numerical accuracy"],
    ],
    [2100, 1750, 2510, 3000],
)
add_body(doc, "The Histogram GB training weight was w = 1 + 5 × award_share. This preserves all zero-share players while giving progressively more influence to serious candidates. Predictions were bounded to [0,1].")

doc.add_page_break()
doc.add_heading("7. Locked final-test results", level=1)
add_table(
    doc,
    ["Model", "RMSE", "Recipient RMSE", "NDCG@5", "Winner rank", "Top-1"],
    [
        ["Histogram GB", "0.0263", "0.1553", "0.8934", "1.25", "75%"],
        ["Extra Trees hurdle", "0.0284", "0.1688", "0.9219", "1.25", "75%"],
        ["Weighted Ridge", "0.0478", "0.2108", "0.8739", "1.25", "75%"],
        ["WS OLS", "0.0583", "0.3460", "0.7195", "2.00", "50%"],
        ["Historical mean", "0.0592", "0.3566", "0.0093", "234.75", "≈0%"],
        ["Zero", "0.0595", "0.3604", "0.0093", "234.75", "≈0%"],
    ],
    [2250, 1250, 1750, 1400, 1400, 1310],
    numeric_columns={1, 2, 3, 4, 5},
)
add_figure(doc, ROOT / "reports" / "figures" / "locked_test_regression.png", "Figure 4. Locked 2019–2022 vote-share error comparison. Lower RMSE is better.")
add_figure(doc, ROOT / "reports" / "figures" / "locked_test_ranking.png", "Figure 5. Locked season-level ranking results. Higher NDCG and lower winner rank are better.")

doc.add_heading("7.1 Winner audit", level=2)
add_table(
    doc,
    ["Season", "Actual MVP", "HGB leader", "Winner rank", "NDCG@5"],
    [
        ["2019", "Giannis Antetokounmpo", "Giannis Antetokounmpo", "1", "0.9590"],
        ["2020", "Giannis Antetokounmpo", "Giannis Antetokounmpo", "1", "0.9284"],
        ["2021", "Nikola Jokić", "Nikola Jokić", "1", "0.8050"],
        ["2022", "Nikola Jokić", "Giannis Antetokounmpo", "2", "0.8812"],
    ],
    [1050, 2550, 2550, 1500, 1710],
    numeric_columns={0, 3, 4},
)
add_figure(doc, ROOT / "reports" / "figures" / "locked_test_winner_rank_heatmap.png", "Figure 6. Actual MVP rank by season for the domain baseline and three learned models.")

doc.add_heading("8. Calibration and imbalance", level=1)
add_body(doc, "The official model’s numerical accuracy is strong, but its boundary behavior requires explicit disclosure. Histogram GB clipped 84.31% of raw predictions, mainly small negative values among zero-share players, and underestimated winning shares by 0.217 on average. The hurdle model required no final clipping and produced the best season-total calibration.")
add_table(
    doc,
    ["Model", "Season-total MAE", "Winner bias", "Clipped"],
    [
        ["Histogram GB", "0.4143", "−0.2170", "84.31%"],
        ["Extra Trees hurdle", "0.1320", "−0.3442", "0%"],
        ["Weighted Ridge", "6.6588", "−0.3277", "62.75%"],
        ["WS OLS", "1.6405", "−0.9031", "32.71%"],
    ],
    [3150, 2150, 1900, 2160],
    numeric_columns={1, 2, 3},
)
add_callout(doc, "Operational rule", "Report HGB and hurdle rankings together. Any disagreement about the top candidate should trigger analyst review rather than silent model selection.", fill="F7F2E8", accent=ORANGE)

doc.add_heading("9. Sensitivity analysis", level=1)
add_body(doc, "The 500-minute sensitivity cohort retained every test vote recipient and all four winners. Histogram GB’s mean winner rank and top-1 accuracy were unchanged, while numerical and ranking metrics moved only modestly. This supports the broader 100-minute cohort as the primary design.")
add_figure(doc, ROOT / "reports" / "figures" / "locked_test_sensitivity.png", "Figure 7. Histogram GB performance under the 100- and 500-minute eligibility thresholds.")

doc.add_page_break()
doc.add_heading("10. Final recommendations", level=1)
add_table(
    doc,
    ["Recommendation", "Implementation"],
    [
        ["Primary forecast", "Use weighted Histogram GB for numerical award_share"],
        ["Ranking check", "Always publish Extra Trees hurdle beside HGB"],
        ["Baseline", "Retain WS percentile OLS in every evaluation"],
        ["Candidate list", "Use mp ≥ 100; retain mp ≥ 500 sensitivity"],
        ["Communication", "Present top five, shares, rank disagreement, and warnings"],
        ["Monitoring", "Track clipping, missingness, drift, NDCG, recipient RMSE, and winner rank"],
        ["Retraining", "Review only after sustained degradation or compatible new seasons"],
    ],
    [2900, 6460],
)

doc.add_heading("11. Limitations", level=1)
add_bullets(doc, [
    "The locked test contains four seasons and only two distinct winners.",
    "A larger untouched modern test window would have produced stronger external evidence.",
    "The target reflects human voting and therefore includes narratives not represented in the table.",
    "The final model is predictive, not causal; it cannot determine who objectively deserves the award.",
    "Modern award-eligibility rules differ from much of the historical period.",
    "No post-2022 dataset was accepted because exact schema compatibility was not verified.",
])
add_body(doc, "A future extension remains appropriate if a candidate source reproduces an overlapping season and preserves the required advanced-statistic definitions. The present conclusions do not depend on such an extension.")

doc.add_heading("12. Conclusion", level=1)
add_body(doc, "Regular-season statistics and team performance can predict MVP vote share with useful accuracy and rank the principal candidates effectively. Histogram GB provided the strongest balance of numerical accuracy and winner identification, while the Extra Trees hurdle model supplied superior top-five ranking and season-total calibration. The combined workflow is most appropriate as transparent decision support, with model disagreement and calibration limitations shown rather than concealed.")
add_callout(doc, "Final answer", "The evidence supports the project hypothesis, but not the stronger claim that MVP voting is completely determined by the available statistics.", fill=PALE_BLUE)

doc.add_page_break()
doc.add_heading("Appendix A. Locked official model", level=1)
add_table(
    doc,
    ["Component", "Locked specification"],
    [
        ["Estimator", "HistGradientBoostingRegressor"],
        ["Learning rate", "0.08"],
        ["Maximum leaf nodes", "31"],
        ["L2 regularization", "1.0"],
        ["Iterations", "200"],
        ["Min. samples per leaf", "20"],
        ["Training weights", "1 + 5 × award_share"],
        ["Final bounds", "Clip predictions to [0,1]"],
        ["Random state", "42"],
    ],
    [3800, 5560],
)

doc.add_heading("Appendix B. Reproducibility assets", level=1)
add_bullets(doc, [
    "final_locked_test_evaluation.py — complete locked evaluation and figures",
    "future_data_schema_audit.py — compatibility gate for prospective datasets",
    "NBA_MVP_PROJECT_HANDOFF.md — reproduction and extension protocol",
    "chunk1_outputs through chunk19_outputs — verified analytical figures",
])

doc.core_properties.title = "Predicting NBA MVP Vote Share"
doc.core_properties.subject = "CRISP-DM regression and season-level ranking analysis"
doc.core_properties.keywords = "NBA, MVP, CRISP-DM, regression, ranking, data science"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""

doc.save(OUT_PATH)
print(OUT_PATH)
